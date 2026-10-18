import streamlit as st
from PIL import Image
import requests
import io
from docx import Document

# Page Configuration
st.set_page_config(
    page_title="AI Herbal Anti-Diabetic Identifier",
    page_icon="🌿",
    layout="centered"
)

# ---------------------------------------------------------
# COMPREHENSIVE ANTI-DIABETIC PLANT DATABASE (WITH PREPARATION GUIDES)
# ---------------------------------------------------------
PLANT_DATABASE = {
    "Vernonia amygdalina": {
        "common_name": "Bitter Leaf",
        "local_names": "Ewuro (Yoruba), Onugbu (Igbo), Shiwaka (Hausa), Etidot (Efik/Ibibio)",
        "has_antidiabetic_properties": True,
        "active_compounds": "Vernodioside, luteolin, saponins, and sesquiterpene lactones.",
        "mechanism": "Improves insulin sensitivity, enhances peripheral glucose uptake, and protects pancreatic beta cells from oxidative stress.",
        "traditional_use": "Fresh leaves are squeezed in clean water to extract the bitter juice, or brewed as a tonic tea.",
        "herb_combinations": "Can be safely combined with African Scent Leaf (Ocimum gratissimum) or Bush Buck/Utazi (Gongronema latifolium) to balance flavor and provide synergistic antioxidant and glucose-lowering effects.",
        "preparation_guide": "• Decoction (Boiling): Boil a handful of clean leaves in 3-4 cups of water for 15-20 minutes. Strain and cool.\n• Cold Infusion (Squeezing): Squeeze fresh leaves in 2 cups of clean drinking water until frothy, then strain.\n• Dosage: Take 1 small glass cup (approx. 150ml) once or twice daily. Store in a cool place and consume within 2-3 days.",
        "safety_warning": "Can significantly lower blood sugar; monitor closely if combined with prescription medications like metformin."
    },
    "Moringa oleifera": {
        "common_name": "Moringa / Drumstick Tree",
        "local_names": "Zogale (Hausa), Okwe Oyibo (Igbo), Ewe Igbale (Yoruba), Mbat (Efik/Ibibio)",
        "has_antidiabetic_properties": True,
        "active_compounds": "Isothiocyanates, quercetin, chlorogenic acid.",
        "mechanism": "Chlorogenic acid helps stabilize post-meal blood sugar levels and assists glucose metabolism.",
        "traditional_use": "Leaves are dried and ground into powder, added to food, or steeped as herbal tea.",
        "herb_combinations": "Pairs exceptionally well with Ginger (Zingiber officinale) or Lemon grass (Cymbopogon citratus) as an infusion to enhance metabolic support and overall anti-inflammatory benefits.",
        "preparation_guide": "• Tea Infusion: Steep 1 tablespoon of leaf powder or dried leaves in 1 cup of boiling water for 10 minutes.\n• Culinary Use: Sprinkle 1 teaspoon of powder into cooked soups or meals daily.\n• Dosage: 1-2 cups of tea daily.",
        "safety_warning": "Safe in culinary amounts; root and bark extracts should be avoided during pregnancy."
    },
    "Ocimum gratissimum": {
        "common_name": "African Scent Leaf / Wild Basil",
        "local_names": "Efinrin (Yoruba), Nchanwu (Igbo), Daidoya (Hausa), Ntong (Efik/Ibibio)",
        "has_antidiabetic_properties": True,
        "active_compounds": "Eugenol, thymol, flavonoids, and polyphenols.",
        "mechanism": "Exhibits strong antioxidant properties and alpha-glucosidase inhibitory activity, slowing carbohydrate digestion.",
        "traditional_use": "Fresh leaves are boiled into a warm tea or infused in water.",
        "herb_combinations": "Frequently combined with Bitter Leaf (Vernonia amygdalina) or Ginger to improve taste and create a balanced daily herbal tonic for blood sugar management.",
        "preparation_guide": "• Decoction: Boil equal parts Scent Leaf and Bitter Leaf in 3 cups of water for 15 minutes.\n• Squeezed Juice: Wash and squeeze leaves in clean water, filter out chaff.\n• Dosage: Take 1 small glass cup warm in the morning before meals.",
        "safety_warning": "Generally safe, but concentrated extracts may interact with blood-thinning medications."
    },
    "Momordica charantia": {
        "common_name": "Wild Bitter Gourd / Balsam Apple",
        "local_names": "Ejirin (Yoruba), Daddagu (Hausa), Nkonnkpaw (Efik/Ibibio)",
        "has_antidiabetic_properties": True,
        "active_compounds": "Charantin, polypeptide-p, and vicine.",
        "mechanism": "Acts similarly to insulin ('plant insulin'), helping transport glucose into body cells.",
        "traditional_use": "Fresh fruit juice or boiled leaf decoction.",
        "herb_combinations": "Can be combined with Fenugreek seeds (Trigonella foenum-graecum) for dual-action glucose lowering, though care must be taken to avoid hypoglycemia.",
        "preparation_guide": "• Decoction: Boil washed leaves or chopped fruit pieces in 3 cups of water for 15-20 minutes.\n• Dosage: Take 1/2 to 1 small glass cup daily. Do not exceed recommended amounts due to high potency.",
        "safety_warning": "Excessive consumption can cause abdominal discomfort or hypoglycemia."
    },
    "Gongronema latifolium": {
        "common_name": "Bush Buck / Bitter Vine",
        "local_names": "Utazi (Igbo), Aroke (Yoruba), Utasi (Efik/Ibibio)",
        "has_antidiabetic_properties": True,
        "active_compounds": "Saponins, tannins, and glycosides.",
        "mechanism": "Improves glucose tolerance and enhances insulin receptor sensitivity.",
        "traditional_use": "Leaves are chewed raw, added to local soups, or brewed into an aqueous extract.",
        "herb_combinations": "Often brewed alongside Bitter Leaf and Scent leaf in traditional Nigerian herbal mixtures to optimize digestive and pancreatic health.",
        "preparation_guide": "• Aqueous Extract: Boil a fresh bundle of Utazi leaves with Scent Leaf in 4 cups of water for 15 minutes.\n• Fresh Consumption: Chew 2-3 fresh washed leaves daily with meals.\n• Dosage: 1 small glass cup daily.",
        "safety_warning": "Use moderately as part of dietary intake if blood sugar is already heavily controlled."
    },
    "Azadirachta indica": {
        "common_name": "Neem Tree",
        "local_names": "Dogoyaro (Hausa/Yoruba), Nim (English)",
        "has_antidiabetic_properties": True,
        "active_compounds": "Azadirachtin, nimbin, nimbidin, and flavonoids.",
        "mechanism": "Helps normalize blood sugar levels by improving insulin receptor sensitivity and reducing oxidative stress.",
        "traditional_use": "Fresh tender leaves or bark are boiled into a bitter decoction and taken as a tonic.",
        "herb_combinations": "Sometimes used in combination with Moringa leaves in mild systemic detox decoctions, but should be used cautiously due to its potency.",
        "preparation_guide": "• Decoction: Boil 5-7 tender fresh leaves in 2 cups of water for 10-15 minutes.\n• Dosage: Take 1/2 small cup every 2-3 days. Use strictly in short durations.",
        "safety_warning": "Excessive use can lower blood pressure and blood sugar too drastically; avoid during pregnancy."
    },
    "Phyllanthus amarus": {
        "common_name": "Stonebreaker / Gale of the Wind",
        "local_names": "Eyin Olobe (Yoruba), Chife-chi (Igbo)",
        "has_antidiabetic_properties": True,
        "active_compounds": "Phyllanthin, hypophyllanthin, flavonoids, and tannins.",
        "mechanism": "Exhibits hypoglycemic effects by improving antioxidant status and protecting pancreatic cells.",
        "traditional_use": "Whole plant decoction or fresh leaf infusion taken orally.",
        "herb_combinations": "Can be paired with Soursop leaves (Annona muricata) for comprehensive renal and metabolic support.",
        "preparation_guide": "• Decoction: Boil clean whole plant (leaves and stems) in 3 cups of water for 15 minutes.\n• Dosage: Take 1 small cup once daily.",
        "safety_warning": "High doses may cause mild gastrointestinal upset; monitor blood glucose levels."
    },
    "Annona muricata": {
        "common_name": "Soursop / Graviola",
        "local_names": "Ewe Aja / Soursop (Yoruba), Uha / Soursop (Igbo)",
        "has_antidiabetic_properties": True,
        "active_compounds": "Annonaceous acetogenins, alkaloids, and flavonoids.",
        "mechanism": "Helps lower blood glucose levels and reduces oxidative stress associated with metabolic complications.",
        "traditional_use": "Fresh leaves are boiled in water to prepare a bitter medicinal tea.",
        "herb_combinations": "Frequently infused alongside Moringa leaves or green tea to provide a soothing daily herbal drink supporting metabolic wellness.",
        "preparation_guide": "• Leaf Tea: Boil 5-10 fresh or dried Soursop leaves in 4 cups of water for 15-20 minutes until the water reduces slightly.\n• Dosage: Take 1 glass cup warm, once daily.",
        "safety_warning": "Long-term excessive use of bark/roots may affect nerve function; consume leaves moderately."
    }
}

# ---------------------------------------------------------
# WORD DOCUMENT GENERATION HELPER
# ---------------------------------------------------------
def generate_word_report(scientific_name, data):
    doc = Document()
    doc.add_heading('🌿 AI Herbal Anti-Diabetic Analysis Report', 0)
    
    doc.add_heading('1. Botanical Identification', level=1)
    doc.add_paragraph(f"Scientific Name: {scientific_name}")
    doc.add_paragraph(f"Common Name: {data['common_name']}")
    doc.add_paragraph(f"Local Names (Nigeria): {data['local_names']}")
    
    doc.add_heading('2. Phytochemical & Pharmacological Profile', level=1)
    doc.add_paragraph(f"Active Compounds: {data['active_compounds']}")
    doc.add_paragraph(f"Mechanism of Action: {data['mechanism']}")
    
    doc.add_heading('3. Traditional Preparation & Usage', level=1)
    doc.add_paragraph(data['traditional_use'])
    
    doc.add_heading('4. Herb Combinations & Synergies', level=1)
    doc.add_paragraph(data['herb_combinations'])

    doc.add_heading('5. Preparation Methods & Dosage Instructions', level=1)
    doc.add_paragraph(data['preparation_guide'])
    
    doc.add_heading('6. Clinical Safety Advisory', level=1)
    doc.add_paragraph(data['safety_warning'])
    
    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer

# ---------------------------------------------------------
# BOTANICAL API RECOGNITION FUNCTION
# ---------------------------------------------------------
def identify_plant_via_api(image_bytes, api_key=""):
    if not api_key or api_key == "YOUR_API_KEY_HERE":
        try:
            api_key = st.secrets.get("PLANTNET_API_KEY", "YOUR_API_KEY_HERE")
        except Exception:
            api_key = "YOUR_API_KEY_HERE"
    
    if not api_key or api_key == "YOUR_API_KEY_HERE":
        return None, "API Key not configured. Using fallback matching mode."
    
    url = f"https://my-api.plantnet.org/v2/identify/all?api-key={api_key}"
    files = [('images', ('plant_image.jpg', image_bytes, 'image/jpeg'))]
    
    try:
        response = requests.post(url, files=files)
        if response.status_code == 200:
            result = response.json()
            if "results" in result and len(result["results"]) > 0:
                best_match = result["results"][0]
                scientific_name = best_match["species"]["scientificNameWithoutAuthor"]
                confidence = best_match["score"]
                return scientific_name, confidence
        return None, "Could not match plant from image."
    except Exception as e:
        return None, str(e)

# ---------------------------------------------------------
# UI INTERFACE
# ---------------------------------------------------------
st.title("🌿 AI Herbal Anti-Diabetic Scanner")
st.markdown("""
Snap or upload a photo of a leaf, bark, or root. The app identifies the botanical species, checks its active compounds, local Nigerian names, traditional combinations, preparation guidelines, and anti-diabetic properties.
""")

with st.sidebar:
    st.header("⚙️ Configuration")
    api_key_input = st.text_input("PlantNet API Key:", type="password")
    if api_key_input:
        st.success("API Key loaded successfully!")
    st.markdown("---")
    st.info("Tip: Enter your PlantNet API key above for live image recognition.")

input_mode = st.radio("Select Input Method:", ("Upload Image", "Snap Photo (Camera)"))

uploaded_file = None
if input_mode == "Upload Image":
    uploaded_file = st.file_uploader("Choose an image file...", type=["jpg", "jpeg", "png"])
else:
    uploaded_file = st.camera_input("Take a picture of the herb, tree, or root")

# ---------------------------------------------------------
# PROCESSING & RESULTS
# ---------------------------------------------------------
if uploaded_file is not None:
    image = Image.open(uploaded_file)
    st.image(image, caption="Captured Specimen", use_container_width=True)
    
    img_byte_arr = io.BytesIO()
    image.save(img_byte_arr, format='JPEG')
    img_bytes = img_byte_arr.getvalue()
    
    st.markdown("---")
    st.subheader("🔍 Analysis Results")
    
    with st.spinner("Processing image and scanning botanical databases..."):
        detected_scientific_name, status = identify_plant_via_api(img_bytes, api_key_input)
        
        if not detected_scientific_name:
            st.warning(f"Notice: {status}. Select a plant to test the output display:")
            detected_scientific_name = st.selectbox(
                "Choose Plant Species:",
                list(PLANT_DATABASE.keys())
            )
            match_confidence = "Simulated Match"
        else:
            match_confidence = f"{status * 100:.1f}% Confidence"
            
    if detected_scientific_name in PLANT_DATABASE:
        plant_info = PLANT_DATABASE[detected_scientific_name]
        
        st.success(f"**Identified Botanical Match:** *{detected_scientific_name}* ({match_confidence})")
        
        st.markdown(f"### **Common Name:** {plant_info['common_name']}")
        st.markdown(f"**Local Names:** {plant_info['local_names']}")
        
        col1, col2 = st.columns(2)
        with col1:
            st.markdown("**Active Phytochemicals:**")
            st.info(plant_info["active_compounds"])
        with col2:
            st.markdown("**Mechanism of Action:**")
            st.info(plant_info["mechanism"])
            
        st.markdown("**Traditional Preparation & Usage:**")
        st.write(plant_info["traditional_use"])
        
        st.warning(f"🌿 **Herb Combinations & Synergies:**\n\n{plant_info['herb_combinations']}")

        st.info(f"🍵 **Preparation Methods & Dosage Guidelines:**\n\n{plant_info['preparation_guide']}")
        
        st.error(f"**Safety Advisory:** {plant_info['safety_warning']}")
        
        st.markdown("---")
        st.subheader("📥 Export Report")
        docx_file = generate_word_report(detected_scientific_name, plant_info)
        st.download_button(
            label="📄 Download Analysis as Word Document (.docx)",
            data=docx_file,
            file_name=f"{detected_scientific_name.replace(' ', '_')}_report.docx",
            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        )
        
    else:
        st.warning(f"The plant *{detected_scientific_name}* was recognized, but it is not currently flagged in our anti-diabetic database.")
        
else:
    st.info("Awaiting image input. Snap or upload a photo to start analyzing.")